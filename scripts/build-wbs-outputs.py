"""Build the two WBS deliverables from the published backlog.

- docs/wbs-master-schedule.md + reports/wbs-master-schedule.xlsx: phases and epics by month
  (Oct 2026 → Jun 2027) with milestones / gates and hours per month.
- docs/wbs-detail.md + reports/wbs-detail.xlsx: epic → issue → sub-issue per sprint with
  role, SWP, hours, confidence tier, status and GitHub link.

Sources: backlog/backlog.json, backlog/swp-hours.json, backlog/github-map.json,
backlog/sprint-s1-applied.json (simulated S1 status). Needs openpyxl.
Usage: <python-with-openpyxl> scripts/build-wbs-outputs.py <repo-root>
"""
import json
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SPRINTS = {  # start, days (calendar) — plan-261024-dev-phase-wbs.md §7
    "P10": ("2026-10-26", 33), "S0": ("2026-11-30", 7), "S1": ("2026-12-07", 14), "S2": ("2026-12-21", 21),
    "S3": ("2027-01-11", 14), "S4": ("2027-01-25", 14), "S5": ("2027-02-08", 21), "S6": ("2027-03-01", 14),
    "HARD": ("2027-03-15", 7), "ST": ("2027-03-22", 14), "UAT": ("2027-04-05", 42), "CUT": ("2027-05-17", 15),
    "HYP": ("2027-06-01", 28),
}
ORDER = list(SPRINTS)
MONTHS = ["2026-10", "2026-11", "2026-12", "2027-01", "2027-02", "2027-03", "2027-04", "2027-05", "2027-06"]
PHASES = [
    ("1.0", "Đóng gap 基本設計 (DR cuối, Step 4/5, ADR Azure, re-estimate)", ["P10"], "G-BD 2026-11-27"),
    ("2.0", "Sprint 0 (repo, CI/CD, IaC, master schedule, backlog, pilot)", ["S0"], "G0 2026-12-04"),
    ("3.0–4.0", "詳細設計 lăn + Implement S1–S6", ["S1", "S2", "S3", "S4", "S5", "S6"], ""),
    ("5.0a", "Hardening (hiệu năng, bảo mật, R-06)", ["HARD"], "G4 2027-03-19"),
    ("5.0b", "System test (IT / ST)", ["ST"], "G5 kỹ thuật 2027-04-02"),
    ("6.0", "移行 (rehearsal 1 ở S6, rehearsal 2 tuần 04-19)", ["S4", "S5", "S6", "UAT"], "Quyết định nạp lần cuối"),
    ("7.0", "UAT + đào tạo 7 vai trò", ["UAT"], "G5 nghiệp vụ 2027-05-14"),
    ("8.0", "Release + cutover (đóng băng master T-5, đơn T-8h)", ["CUT"], "G6 2027-05-31 · Go-Live 2027-06-01"),
    ("9.0", "Chạy song song + hypercare 4 tuần", ["HYP"], "Kết thúc hypercare 2027-06-28"),
    ("10.0", "Quản lý / BrSE (xuyên suốt)", ORDER, "Báo cáo tuần, steering tháng"),
]
HEAD = PatternFill("solid", fgColor="1F4E78")
BAR = PatternFill("solid", fgColor="9BC2E6")
GATE = PatternFill("solid", fgColor="F4B084")


def month_share(sprint):
    """Split a sprint's calendar days across months → {month: fraction}."""
    start, days = SPRINTS[sprint]
    d0 = date.fromisoformat(start)
    share = defaultdict(int)
    for k in range(days):
        share[(d0 + timedelta(days=k)).strftime("%Y-%m")] += 1
    return {m: v / days for m, v in share.items()}


def load(root):
    data = json.loads((root / "backlog" / "backlog.json").read_text(encoding="utf-8"))
    hrs = json.loads((root / "backlog" / "swp-hours.json").read_text(encoding="utf-8"))
    gmap = json.loads((root / "backlog" / "github-map.json").read_text(encoding="utf-8"))["issues"]
    sim = root / "backlog" / "sprint-s1-applied.json"
    status = json.loads(sim.read_text(encoding="utf-8")) if sim.exists() else {}
    return data, hrs, gmap, status


def header(ws, row, values):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_master(root, data, hrs):
    epic_month = defaultdict(lambda: defaultdict(float))
    epic_sprints = defaultdict(set)
    for e in data["epics"]:
        for i in e["issues"]:
            for x in (i.get("sub_issues") or [i]):
                for m, f in month_share(x["sprint"]).items():
                    epic_month[e["id"]][m] += hrs[x["id"]]["hours"] * f
                epic_sprints[e["id"]].add(x["sprint"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Master schedule"
    ws.append(["WBS master schedule — Yuki Cold Chain WMS (giai đoạn phát triển). Milestone theo tháng. Giờ = SWP × 1,85 (docs/hours-convention.md), chưa gồm contingency và phần ngoài backlog."])
    header(ws, 3, ["Mã", "Hạng mục", "Sprint / iteration", "Mốc / gate"] + MONTHS + ["Tổng giờ"])
    r = 4
    for code, name, sps, gate in PHASES:
        months = {m for sp in sps for m in month_share(sp)}
        ws.append([code, name, ", ".join(sps) if len(sps) < 8 else "toàn kỳ", gate] + ["" for _ in MONTHS] + [""])
        for k, m in enumerate(MONTHS):
            if m in months:
                ws.cell(row=r, column=5 + k).fill = GATE if gate and m == max(months) and code != "10.0" else BAR
        ws.cell(row=r, column=1).font = Font(bold=True)
        r += 1
    r += 1
    header(ws, r, ["Epic", "Tên", "Sprint", ""] + MONTHS + ["Tổng giờ"])
    r += 1
    md = ["# WBS master schedule (milestone theo tháng)", "",
          "Giờ = SWP × 1,85 (`docs/hours-convention.md`), chưa gồm contingency và phần ngoài backlog. "
          "Ô = giờ việc rơi vào tháng đó (chia theo số ngày của sprint).", "",
          "## Giai đoạn và mốc", "", "| Mã | Hạng mục | Sprint | Mốc / gate |", "|---|---|---|---|"]
    md += [f"| {c} | {n} | {', '.join(s) if len(s) < 8 else 'toàn kỳ'} | {g} |" for c, n, s, g in PHASES]
    md += ["", "## Giờ theo epic × tháng", "", "| Epic | " + " | ".join(MONTHS) + " | Tổng |", "|---|" + "---:|" * (len(MONTHS) + 1)]
    col_tot = defaultdict(float)
    for e in data["epics"]:
        vals = [round(epic_month[e["id"]].get(m, 0)) for m in MONTHS]
        sps = sorted(epic_sprints[e["id"]], key=ORDER.index)
        ws.append([e["id"], e["title"], f"{sps[0]}–{sps[-1]}", ""] + [v or None for v in vals] + [sum(vals)])
        for k, v in enumerate(vals):
            col_tot[MONTHS[k]] += v
            if v:
                ws.cell(row=r, column=5 + k).fill = BAR
        md.append(f"| {e['id']} {e['title'][:40]} | " + " | ".join(str(v) if v else "" for v in vals) + f" | {sum(vals)} |")
        r += 1
    tot = [round(col_tot[m]) for m in MONTHS]
    ws.append(["", "Tổng", "", ""] + tot + [sum(tot)])
    ws.cell(row=r, column=2).font = Font(bold=True)
    md.append("| **Tổng** | " + " | ".join(f"**{v}**" for v in tot) + f" | **{sum(tot)}** |")
    for c, w in zip("ABCD", (10, 60, 14, 30)):
        ws.column_dimensions[c].width = w
    for k in range(len(MONTHS) + 1):
        ws.column_dimensions[get_column_letter(5 + k)].width = 10
    ws.freeze_panes = "E4"
    wb.save(root / "reports" / "wbs-master-schedule.xlsx")
    (root / "docs" / "wbs-master-schedule.md").write_text("\n".join(md) + "\n", encoding="utf-8")


def build_detail(root, data, hrs, gmap, status):
    wb = Workbook()
    ws = wb.active
    ws.title = "WBS chi tiết"
    cols = ["Mã", "Cấp", "Tên", "Sprint", "Role", "SWP", "Giờ", "Độ tin", "Trạng thái (S1 giả lập)", "GitHub"]
    header(ws, 1, cols)
    md = ["# WBS chi tiết (epic → issue → sub-issue theo sprint)", "",
          "Nguồn: GitHub Project https://github.com/users/linhvtm-2475/projects/1. SWP do `/tkm:sun-estimate` chấm và Linh kiểm mẫu; "
          "giờ = SWP × 1,85. Bảng đầy đủ từng sub-issue ở `reports/wbs-detail.xlsx`; dưới đây là epic → issue.", ""]
    by_sprint = defaultdict(lambda: defaultdict(float))
    for e in data["epics"]:
        ev = hrs[e["id"]]
        ws.append([e["id"], "Epic", e["title"], "", "", ev["swp"], ev["hours"], "", "",
                   f"https://github.com/linhvtm-2475/yuki-wms-backlog/issues/{gmap[e['id']]['number']}"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
        md += [f"## {e['id']} {e['title']} — {ev['swp']} SWP / {ev['hours']:.0f} h", "",
               "| Issue | Sprint | Role | SWP | Giờ | Độ tin | Trạng thái |", "|---|---|---|---:|---:|---|---|"]
        for i in e["issues"]:
            iv = hrs[i["id"]]
            url = f"https://github.com/linhvtm-2475/yuki-wms-backlog/issues/{gmap[i['id']]['number']}"
            st = status.get(i["id"], "Todo")
            ws.append([i["id"], "Issue", i["title"], i["sprint"], i["role"], iv["swp"], iv["hours"],
                       (iv["tier"] or "").replace("estimate:", ""), st, url])
            md.append(f"| [{i['id']}]({url}) {i['title'][:60]} | {i['sprint']} | {i['role']} | {iv['swp']} | {iv['hours']:.0f} | "
                      f"{(iv['tier'] or '').replace('estimate:', '')} | {st} |")
            for s in i.get("sub_issues", []):
                sv = hrs[s["id"]]
                ws.append([s["id"], "Sub-issue", "    " + s["title"], s["sprint"], s["role"], sv["swp"], sv["hours"],
                           (sv["tier"] or "").replace("estimate:", ""), status.get(s["id"], "Todo"),
                           f"https://github.com/linhvtm-2475/yuki-wms-backlog/issues/{gmap[s['id']]['number']}"])
            for x in (i.get("sub_issues") or [i]):
                by_sprint[x["sprint"]][e["id"]] += hrs[x["id"]]["hours"]
        md.append("")
    for c, w in zip("ABCDEFGHIJ", (14, 10, 70, 8, 7, 6, 7, 12, 14, 52)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Giờ theo sprint × epic")
    epics = [e["id"] for e in data["epics"]]
    header(ws2, 1, ["Sprint", "Bắt đầu", "Số ngày"] + epics + ["Tổng"])
    for sp in ORDER:
        vals = [round(by_sprint[sp].get(e, 0)) for e in epics]
        ws2.append([sp, SPRINTS[sp][0], SPRINTS[sp][1]] + [v or None for v in vals] + [sum(vals)])
    ws2.freeze_panes = "B2"
    wb.save(root / "reports" / "wbs-detail.xlsx")
    (root / "docs" / "wbs-detail.md").write_text("\n".join(md) + "\n", encoding="utf-8")


def main(root):
    root = Path(root)
    data, hrs, gmap, status = load(root)
    build_master(root, data, hrs)
    build_detail(root, data, hrs, gmap, status)
    print("wrote docs/wbs-master-schedule.md, docs/wbs-detail.md, reports/wbs-master-schedule.xlsx, reports/wbs-detail.xlsx")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
